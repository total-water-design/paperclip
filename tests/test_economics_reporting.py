import io
import json
import unittest
import zipfile

from openpyxl import load_workbook
from pypdf import PdfReader

from app import app
from economics_reporting import WORKSHEETS, build_pdf, build_snapshot, build_workbook, validate_workbook
from economics import economic_analysis


def model(name="Owner-funded", debt=False, linked=False):
    data = {
        "model": "total_economic_design",
        "project": {"project_name": name, "project_id": f"PRJ-{name[:3].upper()}", "revision_id": "REV-1", "currency": "USD", "capacity_m3d": 10000},
        "scenario_id": "base", "source_summaries": [],
        "cost_items": [{"item_id": "capex-1", "description": "Equipment", "bucket": "equipment_purchase", "quantity": 1, "unit": "LS", "unit_cost": 10000000, "currency": "USD", "source_type": "user", "source_reference": "Estimate"}],
        "allowances": {}, "operating": {"capacity_m3d": 10000, "availability": .95, "fixed_opex_y": 1000000, "other_opex_y": 0, "project_life_years": 25, "discount_rate": .06},
        "finance": {"debt_fraction": .7 if debt else 0, "interest_rate": .06, "debt_tenor_years": 20, "target_dscr": 1.3, "tariff_m3": 1.0},
        "land_rights": [{"id": "ROW-1", "description": "Pipeline corridor", "annual_cost": 50000}] if linked else [],
        "insurance": [{"id": "PRI-1", "category": "expropriation", "premium": 20000}] if debt else [],
    }
    if debt:
        data["project_finance"] = {"enabled": True, "construction_months": 24, "debt_fraction": .7, "construction_debt_rate": .06, "debt_interest_rate": .06, "debt_tenor_years": 20, "concession_years": 25, "repayment_profile": "equal_principal", "tariff_m3": 1, "corporate_tax_rate": .25, "cost_of_equity": .15, "dsra_months": 3}
    if linked:
        data["source_summaries"] = [{"contract": "twds.economic_summary", "version": "1.0", "summary_id": "ro-zld-1", "source": {"application_id": "ro", "application_name": "Total RO Design", "project_id": "RO-1"}, "currency": "USD", "capex": {"buckets": {"equipment_purchase": 1000000}}, "opex": {"annual": {"energy": 100000}}}]
    return data


class EconomicsReportingTests(unittest.TestCase):
    def test_three_required_models_render_both_pdfs_and_valid_workbooks(self):
        for payload in (model(), model("BOOT blended CPI", True), model("TWDS linked RO ZLD", True, True)):
            snapshot = build_snapshot(payload, economic_analysis(payload))
            workbook = build_workbook(snapshot)
            validation = validate_workbook(workbook, snapshot)
            self.assertEqual(validation["worksheets"], list(WORKSHEETS))
            self.assertEqual(validation["formula_cells"], 0)
            self.assertEqual(len(PdfReader(io.BytesIO(build_pdf(snapshot, "executive"))).pages), 17)
            self.assertGreater(len(PdfReader(io.BytesIO(build_pdf(snapshot, "full"))).pages), 30)

    def test_http_exports_share_one_immutable_snapshot(self):
        app.config.update(TESTING=True, AUTH_ENABLED=False)
        with app.test_client() as client:
            created = client.post("/api/economics/report-snapshot", json=model()).get_json()
            self.assertIn("snapshot_id", created)
            frozen = client.get(created["json_url"])
            book = client.get(created["xlsx_url"])
            executive = client.get(created["executive_pdf_url"])
            full = client.get(created["full_pdf_url"])
            self.assertEqual({frozen.status_code, book.status_code, executive.status_code, full.status_code}, {200})
            self.assertEqual(json.loads(frozen.data)["snapshot_id"], created["snapshot_id"])
            wb = load_workbook(io.BytesIO(book.data), read_only=True)
            self.assertEqual(wb["Cover"]["B7"].value, created["snapshot_id"])
            self.assertTrue(executive.data.startswith(b"%PDF-")); self.assertTrue(full.data.startswith(b"%PDF-"))

    def test_security_payloads_are_inert_and_paths_are_not_accepted(self):
        payload = model("=HYPERLINK(\"https://evil.invalid\")")
        payload["cost_items"][0]["description"] = "{{ cycler.__init__.__globals__.os }}"
        snapshot = build_snapshot(payload, economic_analysis(payload))
        data = build_workbook(snapshot)
        wb = load_workbook(io.BytesIO(data), data_only=False)
        self.assertTrue(str(wb["Cover"]["B2"].value).startswith("'="))
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            self.assertFalse(any(x.startswith("xl/externalLinks/") for x in archive.namelist()))
        app.config.update(TESTING=True, AUTH_ENABLED=False)
        with app.test_client() as client:
            self.assertEqual(client.get("/api/economics/report/../../etc/passwd.xlsx").status_code, 404)

    def test_tweco_a_and_c_contract_outputs_route_to_auditable_sheets(self):
        payload = model("Dependency contract")
        payload["economic_model"] = {
            "contract_id": "twds.economic_model", "contract_version": "1.0", "model_id": "MODEL-264",
            "rows": [{"row_id": "ROW-A", "description": "Architecture row", "value": 1}],
            "scenarios": [{"scenario_id": "base", "name": "Base"}],
            "source_lineage": [{"lineage_id": "SRC-A", "source_type": "quotation", "reference": "Q-100"}],
            "audit_events": [{"event_id": "AUD-A", "action": "create"}],
        }
        result = economic_analysis(payload)
        result["capex_wbs"] = {
            "nodes": [{"wbs_id": "WBS-C", "name": "Imported equipment", "amount": 250000}],
            "rollups": [{"wbs_id": "WBS-C", "amount": 250000}], "total": 250000,
            "included_total": 250000, "excluded_total": 0,
            "procurement": [{"wbs_id": "WBS-C", "po_date": "2026-09-01"}],
            "logistics": [{"wbs_id": "WBS-C", "arrival_date": "2027-01-01"}],
            "construction_schedule": {"monthly_spend": [{"month": "2026-09", "amount": 250000}]},
        }
        snapshot = build_snapshot(payload, result)
        self.assertEqual(snapshot["model_id"], "MODEL-264")
        wb = load_workbook(io.BytesIO(build_workbook(snapshot)), read_only=True)
        sheet_text = lambda name: " ".join(str(cell.value) for row in wb[name].iter_rows() for cell in row if cell.value is not None)
        self.assertIn("WBS-C", sheet_text("WBS"))
        self.assertIn("Q-100", sheet_text("Connected Sources"))
        self.assertIn("2027-01-01", sheet_text("Procurement & Logistics"))
        self.assertIn("2026-09", sheet_text("Construction"))
        self.assertIn("AUD-A", sheet_text("Audit Trail"))


if __name__ == "__main__": unittest.main()
