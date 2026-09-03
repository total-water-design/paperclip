"""Immutable, auditable reporting exports for Total Water Economics."""
from __future__ import annotations

import hashlib
import io
import json
import math
import re
import uuid
import zipfile
from copy import deepcopy
from datetime import datetime, timezone
from xml.sax.saxutils import escape

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from pypdf import PdfReader
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

SNAPSHOT_CONTRACT = "twds.economics_report_snapshot"
SNAPSHOT_VERSION = "1.0"
WORKSHEETS = (
    "Cover", "Model Checks", "Assumptions", "Connected Sources", "CPI Data",
    "Inflation Profiles", "WBS", "CAPEX", "Procurement & Logistics", "Land & ROW",
    "Insurance", "Construction", "Sources & Uses", "OPEX", "Revenue", "Debt",
    "Reserves", "Tax", "Cash Flow", "Returns", "Scenarios", "Sensitivity", "Audit Trail",
)


def _canonical(value: object) -> bytes:
    return json.dumps(value, sort_keys=True, separators=(",", ":"), ensure_ascii=False, allow_nan=False).encode()


def _safe_text(value: object, limit: int = 32000) -> str:
    text = str("" if value is None else value).replace("\x00", "").strip()[:limit]
    # Prevent spreadsheet formula/DDE execution while preserving the visible value.
    return "'" + text if text.startswith(("=", "+", "-", "@")) else text


def _finite(value: object) -> bool:
    return not isinstance(value, float) or math.isfinite(value)


def _walk(value: object, path: str = "$"):
    if isinstance(value, dict):
        for key, child in value.items():
            yield from _walk(child, f"{path}.{key}")
    elif isinstance(value, list):
        for index, child in enumerate(value):
            yield from _walk(child, f"{path}[{index}]")
    else:
        yield path, value


def build_snapshot(model_input: dict, result: dict) -> dict:
    if not isinstance(model_input, dict) or not isinstance(result, dict):
        raise ValueError("Model input and calculated result must be objects.")
    frozen_input, frozen_result = deepcopy(model_input), deepcopy(result)
    for path, value in _walk({"input": frozen_input, "result": frozen_result}):
        if not _finite(value):
            raise ValueError(f"Non-finite report value at {path}.")
    project = frozen_input.get("project") or {}
    project_id = _safe_text(project.get("project_id") or frozen_input.get("project_id") or "UNASSIGNED", 120)
    revision_id = _safe_text(project.get("revision_id") or frozen_input.get("revision_id") or "UNREVISIONED", 120)
    scenario_id = _safe_text(frozen_input.get("scenario_id") or "base", 120)
    model_document = frozen_input.get("economic_model") or frozen_input.get("model_document") or {}
    model_id = _safe_text(model_document.get("model_id") or frozen_input.get("model_id") or "UNASSIGNED", 120)
    calculation_hash = hashlib.sha256(_canonical(frozen_result)).hexdigest()
    snapshot_hash = hashlib.sha256(_canonical({"input": frozen_input, "result": frozen_result})).hexdigest()
    snapshot_id = str(uuid.uuid5(uuid.NAMESPACE_URL, snapshot_hash))
    snapshot = {
        "contract": SNAPSHOT_CONTRACT, "version": SNAPSHOT_VERSION,
        "snapshot_id": snapshot_id, "calculation_id": calculation_hash,
        "created_at": datetime.now(timezone.utc).isoformat(),
        "project_id": project_id, "revision_id": revision_id, "scenario_id": scenario_id,
        "model_id": model_id,
        "project_name": _safe_text(project.get("project_name") or "Untitled Water Economics Project", 240),
        "reporting_currency": _safe_text(project.get("currency") or "USD", 12),
        "model_input": frozen_input, "calculation_result": frozen_result,
        "headlines": analytical_headlines(frozen_result),
        "limitations": [
            "TWDS is the authoritative calculation engine; exported workbooks are auditable outputs.",
            "Results depend on the supplied assumptions, source vintages and scenario definition.",
            "No exported narrative constitutes financing, legal, tax or insurance advice.",
        ],
    }
    validate_snapshot(snapshot)
    return json.loads(_canonical(snapshot))


def validate_snapshot(snapshot: dict) -> None:
    required = ("contract", "version", "snapshot_id", "calculation_id", "project_id", "revision_id", "scenario_id", "model_id", "model_input", "calculation_result")
    missing = [key for key in required if key not in snapshot]
    if missing:
        raise ValueError("Report snapshot missing: " + ", ".join(missing))
    if snapshot["contract"] != SNAPSHOT_CONTRACT or snapshot["version"] != SNAPSHOT_VERSION:
        raise ValueError("Unsupported economics report snapshot contract.")
    for path, value in _walk(snapshot):
        if not _finite(value):
            raise ValueError(f"Non-finite report value at {path}.")


def _metric(result: dict, *paths: str):
    for path in paths:
        value = result
        for key in path.split("."):
            value = value.get(key) if isinstance(value, dict) else None
        if isinstance(value, (int, float)) and _finite(value):
            return value
    return None


def analytical_headlines(result: dict) -> list[str]:
    currency = result.get("reporting_currency") or "USD"
    total = _metric(result, "cost_hierarchy.total_project_cost", "total_project_cost", "total_capex")
    lcow = _metric(result, "lifecycle.lcow", "lcow", "finance.lcow")
    dscr = _metric(result, "project_finance.debt.min_dscr", "finance.min_dscr")
    required_tariff = _metric(result, "project_finance.returns.required_tariff", "finance.required_tariff")
    items = []
    if total is not None: items.append(f"Total project cost is {currency} {total:,.0f} on the frozen calculation basis")
    if lcow is not None: items.append(f"Levelized water cost is {lcow:,.3f} {currency}/m³")
    if dscr is not None: items.append(f"Minimum debt-service coverage is {dscr:,.2f}× over the modeled tenor")
    if required_tariff is not None: items.append(f"Required modeled tariff is {required_tariff:,.3f} {currency}/m³")
    warnings = result.get("warnings") or (result.get("project_finance") or {}).get("warnings") or []
    if warnings: items.append(f"Model validation identifies {len(warnings)} item(s) requiring review")
    return items[:5] or ["The immutable calculation snapshot contains no recognized headline metrics"]


def snapshot_json(snapshot: dict) -> bytes:
    validate_snapshot(snapshot)
    return json.dumps(snapshot, indent=2, ensure_ascii=False, allow_nan=False, sort_keys=True).encode()


def _rows(value: object, prefix: str = "") -> list[list[object]]:
    rows = []
    if isinstance(value, dict):
        for key, child in value.items(): rows.extend(_rows(child, f"{prefix}.{key}" if prefix else str(key)))
    elif isinstance(value, list):
        for index, child in enumerate(value): rows.extend(_rows(child, f"{prefix}[{index}]"))
    else: rows.append([prefix, value])
    return rows


def build_workbook(snapshot: dict) -> bytes:
    validate_snapshot(snapshot)
    wb = Workbook(); wb.remove(wb.active)
    inputs, result = snapshot["model_input"], snapshot["calculation_result"]
    model_document = inputs.get("economic_model") or inputs.get("model_document") or {}
    wbs = result.get("capex_wbs") or inputs.get("capex_wbs") or {}
    connected_sources = list(inputs.get("source_summaries") or []) + list(model_document.get("source_lineage") or [])
    sections = {
        "Assumptions": model_document.get("rows") or inputs, "Connected Sources": connected_sources,
        "CPI Data": inputs.get("cpi_data", []), "Inflation Profiles": inputs.get("inflation_profiles", []),
        "WBS": wbs.get("nodes") or inputs.get("wbs", inputs.get("cost_items", [])), "CAPEX": {"cost_hierarchy": result.get("cost_hierarchy", {}), "wbs_rollups": wbs.get("rollups", []), "wbs_totals": {key: wbs.get(key) for key in ("total", "included_total", "excluded_total", "by_scope", "by_sourcing") if key in wbs}},
        "Procurement & Logistics": {"procurement": wbs.get("procurement", inputs.get("procurement", [])), "logistics": wbs.get("logistics", inputs.get("logistics", []))}, "Land & ROW": inputs.get("land_rights", []),
        "Insurance": inputs.get("insurance", []), "Construction": (result.get("project_finance") or {}).get("construction", {}),
        "Sources & Uses": (result.get("project_finance") or {}).get("sources_and_uses", {}),
        "OPEX": result.get("opex", inputs.get("operating", {})), "Revenue": (result.get("project_finance") or {}).get("revenue", {}),
        "Debt": (result.get("project_finance") or {}).get("debt", {}), "Reserves": (result.get("project_finance") or {}).get("reserves", {}),
        "Tax": (result.get("project_finance") or {}).get("tax", {}), "Cash Flow": (result.get("project_finance") or {}).get("operations", {}),
        "Returns": (result.get("project_finance") or {}).get("returns", result.get("lifecycle", {})),
        "Scenarios": model_document.get("scenarios") or inputs.get("scenarios", []), "Sensitivity": inputs.get("sensitivity", []), "Audit Trail": model_document.get("audit_events") or inputs.get("audit_trail", []),
    }
    if wbs.get("construction_schedule"):
        sections["Construction"] = {"project_finance": sections["Construction"], "capex_wbs": wbs["construction_schedule"]}
    for title in WORKSHEETS:
        ws = wb.create_sheet(title); ws.freeze_panes = "A2"; ws.sheet_view.showGridLines = False
        ws.append(["Total Water Economics", title]); ws["A1"].font = Font(bold=True, color="FFFFFF"); ws["B1"].font = Font(bold=True, color="FFFFFF")
        for cell in ws[1]: cell.fill = PatternFill("solid", fgColor="985A00")
        if title == "Cover":
            rows = [["Project", snapshot["project_name"]], ["Project ID", snapshot["project_id"]], ["Revision ID", snapshot["revision_id"]], ["Scenario ID", snapshot["scenario_id"]], ["Calculation ID", snapshot["calculation_id"]], ["Report snapshot ID", snapshot["snapshot_id"]], ["Model ID", snapshot["model_id"]], ["Created UTC", snapshot["created_at"]], ["Currency", snapshot["reporting_currency"]], ["Authority", "TWDS calculation engine"]]
        elif title == "Model Checks":
            rows = [["Snapshot contract", "PASS"], ["Finite values", "PASS"], ["Immutable calculation hash", snapshot["calculation_id"]]] + [["Headline", x] for x in snapshot["headlines"]] + [["Limitation", x] for x in snapshot["limitations"]]
        else: rows = _rows(sections.get(title, {})) or [["status", "No data supplied in snapshot"]]
        for key, value in rows:
            ws.append([_safe_text(key), value if isinstance(value, (int, float, bool)) else _safe_text(value)])
        ws.column_dimensions["A"].width = 46; ws.column_dimensions["B"].width = 90
        for row in ws.iter_rows():
            for cell in row: cell.alignment = Alignment(vertical="top", wrap_text=True)
    wb.calculation.fullCalcOnLoad = False; wb.calculation.forceFullCalc = False
    out = io.BytesIO(); wb.save(out); validate_workbook(out.getvalue(), snapshot); return out.getvalue()


def validate_workbook(data: bytes, snapshot: dict | None = None) -> dict:
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        names = archive.namelist()
        if any(name.lower().endswith(("vbaproject.bin", ".exe", ".js")) for name in names): raise ValueError("Executable workbook content is prohibited.")
        external = [name for name in names if name.startswith("xl/externalLinks/")]
        if external: raise ValueError("External workbook links are prohibited.")
    wb = load_workbook(io.BytesIO(data), data_only=False, read_only=True)
    missing = [name for name in WORKSHEETS if name not in wb.sheetnames]
    if missing: raise ValueError("Workbook missing sheets: " + ", ".join(missing))
    bad = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.data_type == "f" or (isinstance(cell.value, str) and re.search(r"#(?:REF!|DIV/0!|VALUE!|NAME\?)", cell.value)): bad.append(f"{ws.title}!{cell.coordinate}")
    if bad: raise ValueError("Workbook contains formulas/errors: " + ", ".join(bad[:10]))
    if snapshot and wb["Cover"]["B7"].value != snapshot["snapshot_id"]: raise ValueError("Workbook snapshot ID does not reconcile.")
    return {"worksheets": wb.sheetnames, "formula_cells": 0, "external_links": 0, "macros": False}


EXECUTIVE_SECTIONS = ("Executive Investment Summary", "Project and Transaction Overview", "Investment Case", "Capital Cost and WBS", "Inflation and Indexation", "Procurement and Logistics", "Land and Rights-of-Way", "Insurance and Guarantees", "Sources & Uses", "Construction and Funding", "Operating Economics", "Revenue and Tariff", "Project Finance", "Cash Flow and Returns", "Sensitivity and Scenarios", "Risks, Assumptions and Provenance")


def build_pdf(snapshot: dict, report_type: str) -> bytes:
    validate_snapshot(snapshot)
    if report_type not in {"executive", "full"}: raise ValueError("Report type must be executive or full.")
    out = io.BytesIO(); styles = getSampleStyleSheet()
    doc = SimpleDocTemplate(out, pagesize=A4 if report_type == "executive" else landscape(A4), rightMargin=16*mm, leftMargin=16*mm, topMargin=16*mm, bottomMargin=14*mm, title=f"{snapshot['project_name']} - {report_type.title()} Report", author="Total Water Design Suite")
    story = [Paragraph("TOTAL WATER DESIGN SUITE", styles["Title"]), Spacer(1, 8), Paragraph("Total Water Economics", styles["Heading1"]), Paragraph(escape(snapshot["project_name"]), styles["Heading2"]), Spacer(1, 16)]
    meta = [["Project / revision", f"{snapshot['project_id']} / {snapshot['revision_id']}"], ["Model / scenario", f"{snapshot['model_id']} / {snapshot['scenario_id']}"], ["Calculation snapshot", snapshot["snapshot_id"]], ["Calculation ID", snapshot["calculation_id"]], ["Design status", "CONFIDENTIAL — MODEL OUTPUT"]]
    story += [Table(meta, colWidths=[42*mm, 125*mm]), PageBreak()]
    sections = list(EXECUTIVE_SECTIONS)
    if report_type == "full": sections += [name for name in WORKSHEETS[2:] if name not in sections]
    flat_result = _rows(snapshot["calculation_result"])
    for index, title in enumerate(sections):
        story.append(Paragraph(escape(title), styles["Heading1"]))
        story.append(Paragraph(escape(snapshot["headlines"][index % len(snapshot["headlines"])]), styles["Heading2"]))
        source = flat_result[index * 8:(index + 1) * 8] or [["Snapshot", snapshot["snapshot_id"]]]
        table = Table([["Metric / source path", "Frozen value"]] + [[_safe_text(k, 120), _safe_text(v, 240)] for k, v in source], colWidths=[75*mm, 100*mm], repeatRows=1)
        table.setStyle(TableStyle([("BACKGROUND", (0,0), (-1,0), colors.HexColor("#985A00")), ("TEXTCOLOR", (0,0), (-1,0), colors.white), ("GRID", (0,0), (-1,-1), .25, colors.HexColor("#B9B9B9")), ("VALIGN", (0,0), (-1,-1), "TOP"), ("FONTSIZE", (0,0), (-1,-1), 7)]))
        story += [table, Spacer(1, 8), Paragraph("Source: immutable TWDS calculation snapshot. Limitations apply.", styles["BodyText"])]
        if index != len(sections)-1: story.append(PageBreak())
    doc.build(story)
    data = out.getvalue()
    reader = PdfReader(io.BytesIO(data))
    expected_pages = 1 + len(sections)
    if not data.startswith(b"%PDF-") or len(reader.pages) != expected_pages: raise ValueError("PDF validation failed.")
    cover_text = reader.pages[0].extract_text() or ""
    if snapshot["snapshot_id"] not in cover_text or snapshot["project_id"] not in cover_text: raise ValueError("PDF identifiers do not reconcile.")
    return data
